#!/usr/bin/env python3
"""Versión imprimible (.docx) de un escrito en Markdown, con la identidad de Poïesis.

Uso:
  python3 herramientas/escrito_docx.py ESCRITO.md --caso casos/<caso> [--salida RUTA.docx] [--pdf] [--confidencial]

Retira las citas [R-xxx]/[N-xxx] y los comentarios <!-- -->. Antes de generar corre
el cotejo del escrito: si hay observaciones, el documento sale marcado como BORRADOR
en el encabezado de cada página. Admite el subconjunto de Markdown que usan los
escritos: títulos # y ##, párrafos, **negrita**, listas con - y 1., tablas y ---
(salto de página). El título # va centrado, y también el subtítulo en negrita que lo sigue.
"""
import argparse
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
import registro  # noqa: E402

NAVY, TINTA, GRIS, ARENA = "131237", "2F2E47", "5C6684", "D1D0CB"
PALETA = (NAVY, TINTA, GRIS, ARENA, "FFFFFF")
COLOR = re.compile(r'((?:w:color w:val|w:fill|a:srgbClr val|lastClr)=")([0-9A-Fa-f]{6})(")')
TITULOS, TEXTO = "Safira March", "Georgia"
CITA = re.compile(r"\s*\[(?:[RN]-\d{3})(?:\s*,\s*[RN]-\d{3})*\]")
NEGRITA = re.compile(r"\*\*(.+?)\*\*")


def color(hexa):
    return RGBColor.from_string(hexa)


def cercano(hexa):
    rgb = [int(hexa[i:i + 2], 16) for i in (0, 2, 4)]
    return min(PALETA, key=lambda p: sum((int(p[i:i + 2], 16) - v) ** 2 for i, v in zip((0, 2, 4), rgb)))


def a_paleta(ruta):
    """Lleva todo color del archivo, estilos y tema de Word incluidos, al tono más cercano de la paleta."""
    with zipfile.ZipFile(ruta) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        for n, datos in partes.items():
            if n.endswith(".xml"):
                texto = COLOR.sub(lambda m: m.group(1) + cercano(m.group(2)) + m.group(3), datos.decode("utf-8"))
                datos = texto.encode("utf-8")
            z.writestr(n, datos)


def runs(p, texto, size=10.5, tinta=TINTA, negrita=False):
    """Agrega el texto al párrafo respetando **negrita** en línea."""
    partes = NEGRITA.split(texto)
    for i, parte in enumerate(partes):
        if not parte:
            continue
        r = p.add_run(parte)
        r.font.name, r.font.size, r.font.color.rgb = TEXTO, Pt(size), color(tinta)
        r.bold = negrita or i % 2 == 1
    return p


def espacio(p, antes=0, despues=6, interlineado=1.25):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(antes), Pt(despues), interlineado


def borde_inferior(p, hexa=ARENA):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", hexa)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def sombrear(celda, hexa):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", hexa)):
        shd.set(qn(k), v)
    tcpr.append(shd)


def bordes(tabla, visibles):
    tblpr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        if visibles:
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", ARENA)):
                e.set(qn(k), v)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def campo(p, instr):
    for tipo in ("begin", None, "end"):
        r = p.add_run()
        r.font.size, r.font.color.rgb = Pt(8.5), color(GRIS)
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = instr
        r._r.append(e)


def tabla(doc, filas, ancho):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    cab, cuerpo = celdas[0], [f for f in celdas[2:]]
    con_cabecera = any(cab)
    datos = ([cab] if con_cabecera else []) + cuerpo
    n = max(len(f) for f in datos)
    firmas = datos[0][0].startswith("_")
    t = doc.add_table(rows=len(datos), cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes(t, con_cabecera)
    for i, fila in enumerate(datos):
        for j in range(n):
            c = t.cell(i, j)
            c.width = Inches(ancho / n)
            p = c.paragraphs[0]
            espacio(p, 2, 2, 1.1)
            texto = fila[j] if j < len(fila) else ""
            if con_cabecera and i == 0:
                sombrear(c, ARENA)
                runs(p, texto, 9.5, NAVY, True)
            else:
                etiqueta = not con_cabecera and not firmas and n == 2 and j == 0
                runs(p, texto, 9.5, GRIS if etiqueta else TINTA)
    doc.add_paragraph()


def convertir(md, salida, aviso, confidencial=False):
    texto = re.sub(r"<!--.*?-->", "", md, flags=re.S)
    texto = CITA.sub("", texto)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.color.rgb = TEXTO, color(TINTA)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Inches(1.1))
    sec.top_margin, sec.bottom_margin = Inches(1), Inches(0.9)
    ancho = 8.5 - 2.2

    cab = sec.header.paragraphs[0]
    runs(cab, "Grupo Poïesis · Confidencial" if confidencial else "Grupo Poïesis", 8.5, GRIS)
    if aviso:
        cab.add_run("    ")
        runs(cab, aviso, 8.5, NAVY, True)
    pie = sec.footer.paragraphs[0]
    runs(pie, "Página ", 8.5, GRIS)
    campo(pie, "PAGE")
    runs(pie, " de ", 8.5, GRIS)
    campo(pie, "NUMPAGES")

    lineas = texto.splitlines()
    i = 0
    salto = False
    while i < len(lineas):
        l = lineas[i].rstrip()
        if not l.strip():
            i += 1
            continue
        if l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(lineas[i])
                i += 1
            tabla(doc, filas, ancho)
            continue
        if l.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(l[2:])
            r.font.name, r.font.size, r.font.color.rgb = TITULOS, Pt(24), color(NAVY)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            espacio(p, 0, 4)
        elif l.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(l[3:])
            r.font.name, r.font.size, r.font.color.rgb = TITULOS, Pt(14), color(NAVY)
            espacio(p, 14, 6)
            p.paragraph_format.keep_with_next = True
            borde_inferior(p)
        elif l.strip() == "---":
            salto = True
        elif l.strip() == "&nbsp;":
            doc.add_paragraph()
        elif l.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            runs(p, l[2:])
            espacio(p, 0, 3)
        elif re.match(r"^\d+\. ", l):
            p = doc.add_paragraph(style="List Number")
            runs(p, re.sub(r"^\d+\. ", "", l))
            espacio(p, 0, 3)
        elif l.startswith("*") and l.endswith("*") and not l.startswith("**"):
            p = doc.add_paragraph()
            runs(p, l.strip("*"), 10, GRIS)
            for r in p.runs:
                r.italic = True
            espacio(p, 0, 8)
        else:
            p = doc.add_paragraph()
            solo_negrita = l.startswith("**") and l.endswith("**") and l.count("**") == 2
            runs(p, l, 11 if solo_negrita else 10.5, NAVY if solo_negrita else TINTA)
            espacio(p, 6 if solo_negrita else 0, 6)
            if solo_negrita and doc.paragraphs[-2].alignment == WD_ALIGN_PARAGRAPH.CENTER:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if l.startswith("Firmamos"):
            # La firma no queda sola en una página: viaja con la última cláusula.
            for q in doc.paragraphs[-4:]:
                q.paragraph_format.keep_with_next = True
        elif l.strip() == "&nbsp;" and doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        if salto and l.strip() != "---":
            doc.paragraphs[-1].paragraph_format.page_break_before = True
            salto = False
        i += 1
    doc.save(salida)
    a_paleta(salida)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("escrito")
    ap.add_argument("--caso", required=True)
    ap.add_argument("--salida")
    ap.add_argument("--confidencial", action="store_true", help="marca «Confidencial» en el encabezado")
    ap.add_argument("--pdf", action="store_true", help="genera también la versión PDF")
    a = ap.parse_args()
    _, problemas = registro.cotejo(registro.cargar(a.caso), a.escrito)
    aviso = (f"BORRADOR · No entregar: el cotejo tiene {len(problemas)} observaciones"
             if problemas else "")
    salida = a.salida or str(Path(a.escrito).with_suffix(".docx"))
    convertir(Path(a.escrito).read_text(encoding="utf-8"), salida, aviso, a.confidencial)
    print(f"Generado {salida}" + (f" ({aviso})." if aviso else " (cotejo sin observaciones)."))
    if a.pdf:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(Path(salida).parent), salida],
                       check=True, capture_output=True)
        print(f"Generado {Path(salida).with_suffix('.pdf')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
